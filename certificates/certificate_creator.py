from docx import Document
from datetime import datetime
from docx2pdf import convert
import locale
from pathlib import Path
from project.utils import paths

locale.setlocale(locale.LC_TIME, 'pt_BR.UTF-8')

TEMPLATE_DIR = paths.CERTIFICATES_DIR / 'templates' / 'brigada_de_incendio.docx'

class Certificate:
    def __init__(self, template_dir: str, output_dir: str) -> None:
        self.template_dir = Path(template_dir)
        self.output_dir = Path(output_dir)
        self.docx_dir = self.output_dir / 'docx'
        self.pdf_dir = self.output_dir / 'pdf'
        self.docx: Document | None = None
        self._created_date: str = datetime.now().strftime("%d de %B de %Y")

    def __repr__(self) -> str:
        return f'Certificate(template_dir={self.template_dir}, output_dir={self.output_dir})'

    def _ensure_directories(self) -> None:
        # """Ensure that the necessary output directories exist."""
        directories = [self.output_dir, self.docx_dir, self.pdf_dir]
        for directory in directories:
            if not directory.exists():
                directory = Folder(folder_path=directory)
        # self.output_dir.mkdir(exist_ok=True)
        # self.docx_dir.mkdir(exist_ok=True)
        # self.pdf_dir.mkdir(exist_ok=True)

    def _load_template(self) -> None:
        """Load the Word document template."""
        self.docx = Document(self.template_dir)

    def _replace_placeholders(self, replacements: dict[str, str]) -> None:
        """Replace placeholders in the document with actual values."""
        for paragraph in self.docx.paragraphs:
            if any(placeholder in paragraph.text for placeholder in replacements.keys()):
                for run in paragraph.runs:
                    for placeholder, value in replacements.items():
                        run.text = run.text.replace(placeholder, value)

    def _save_document(self, name: str) -> tuple[Path, Path]:
        """Save the document as both DOCX and PDF."""
        name_file = f'Certificado de {name.title()}'
        docx_file = self.docx_dir / f'{name_file}.docx'
        pdf_file = self.pdf_dir / f'{name_file}.pdf'
        self.docx.save(docx_file)
        convert(docx_file, pdf_file)
        return docx_file, pdf_file

    def create_certificate(
        self,
        name: str,
        date: str,
        workload: str,
        company: str,
    ) -> tuple[Path, Path]:
        """
        Create a certificate by replacing placeholders in the template
        and saving it as both DOCX and PDF.
        """
        self._ensure_directories()
        self._load_template()
        replacements = {
            'REPLACE_name': name.title(),
            'REPLACE_date': date,
            'REPLACE_workload': workload,
            'REPLACE_company': company.upper(),
            'REPLACE_current_date': self._created_date,
        }
        self._replace_placeholders(replacements)
        return self._save_document(name)
    

class Folder:
    def __init__(self, folder_path: str) -> None:
        self.folder_path = Path(folder_path)
        self._ensure_directory()

    def _ensure_directory(self) -> None:
        """Ensure that the directory exists."""
        self.folder_path.mkdir(exist_ok=True)

    def __repr__(self) -> str:
        return f'Folder(folder_path={self.folder_path})'
